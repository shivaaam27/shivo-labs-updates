"""Sheet artifact (JSON) -> Excel file. Usage: python sheet2xlsx.py in.json out.xlsx

{"sheets": [{"name": "Budget", "header": true, "rows": [["Item", "Qty"], ["Pens", 10], ["Total", "=SUM(B2:B2)"]]}]}
Values go in as numbers or text; strings starting with = go in as real Excel formulas."""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def main(src, out):
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    wb = Workbook()
    wb.remove(wb.active)
    for i, s in enumerate(data.get("sheets") or [{"name": "Sheet1", "rows": []}]):
        name = "".join(ch for ch in str(s.get("name") or f"Sheet{i + 1}") if ch not in '[]:*?/\\')[:31] or f"Sheet{i + 1}"
        ws = wb.create_sheet(name)
        rows = s.get("rows") or []
        widths = {}
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row, 1):
                if v is None or v == "":
                    continue
                ws.cell(row=r, column=c, value=v)
                shown = len(str(v)) if not (isinstance(v, str) and v.startswith("=")) else 10
                widths[c] = max(widths.get(c, 8), min(shown + 2, 60))
        first = rows[0] if rows else []
        if s.get("header", True) and first and all(isinstance(v, str) and v and not v.startswith("=") for v in first if v not in (None, "")):
            for c in range(1, len(first) + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="F2F2EF")
            ws.freeze_panes = "A2"
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = w
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
